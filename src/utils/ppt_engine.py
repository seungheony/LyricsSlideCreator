from pptx import Presentation
import copy
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from lxml import etree
import io
import os
import re
import datetime
import unicodedata


def find_lyrics_shape(slide):
    """가사를 넣을 글상자: '가사'라고 적힌 첫 글상자(글자가 있는 것). 없으면 None.
    줄바꿈 미리 계산(line_fit)도 같은 글상자를 기준으로 한다."""
    for shape in slide.shapes:
        if shape.has_text_frame and '가사' in shape.text and shape.text_frame.paragraphs[0].runs:
            return shape
    return None


_NEWLINES = re.compile('[\x0b\x0c\u2028\u2029\x85]')
_ZERO_WIDTH = re.compile('[\u200b-\u200d\u2060\ufeff\u00ad]')
# 줄바꿈(\n)을 뺀 나머지 제어 문자, 서러게이트, 문자가 아닌 코드(U+FFFE·U+FFFF)
_CONTROL = re.compile('[\x00-\x09\x0e-\x1f\x7f-\x9f\ud800-\udfff\ufffe\uffff]')


def clean_lyrics_text(text):
    """PPT에 넣기 전 가사 정리: \\v·\\f·U+2028/2029는 줄바꿈으로, 제로폭 문자는 제거, 탭은 공백으로,
    그 밖의 제어 문자는 제거(python-pptx는 제어 문자를 '_x000B_' 같은 글자로 바꿔 찍는다)."""
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = _NEWLINES.sub('\n', text).replace('\t', ' ')
    text = _ZERO_WIDTH.sub('', text)
    return unicodedata.normalize('NFC', _CONTROL.sub('', text))


class PPTBuilder:
    def __init__(self, template_path):
        self.template_path = template_path
        self.prs = None
        try:
            self.prs = Presentation(template_path)
        except Exception as e:
            print(f"Failed to load template: {e}")
            raise

    def duplicate_slide_elements(self, source_slide, target_slide):
        for shape in source_slide.shapes:
            el = shape.element
            new_el = copy.deepcopy(el)
            target_slide.shapes._spTree.insert_element_before(new_el, 'p:extLst')

    def replace_text_in_slide(self, slide, lyrics_chunk):
        shape = find_lyrics_shape(slide)
        if shape is None:
            return

        text_frame = shape.text_frame
        original_p = text_frame.paragraphs[0]
        original_run = original_p.runs[0]

        # 문단 서식(줄 간격·정렬 등)과 글자 서식(폰트·크기·색·언어 등)을 XML째로 보관
        original_pPr = copy.deepcopy(original_p._p.pPr)
        original_rPr = copy.deepcopy(original_run._r.rPr)

        # 첫 번째 문단을 제외한 나머지 문단을 삭제
        for p in text_frame.paragraphs[1:]:
            p._p.getparent().remove(p._p)

        # 첫 번째 문단의 내용을 모두 지움 (문단 서식은 유지됨)
        p = text_frame.paragraphs[0]
        p.clear()

        lines = lyrics_chunk.split('\n')
        for i, line in enumerate(lines):
            if i > 0:
                p = text_frame.add_paragraph()
                # 새 문단에도 첫 문단의 서식을 그대로 적용 (없으면 줄 간격이 기본값으로 바뀜)
                if original_pPr is not None:
                    p._p.remove(p._p.get_or_add_pPr())
                    p._p.insert(0, copy.deepcopy(original_pPr))

            run = p.add_run()
            run.text = line
            if original_rPr is not None:
                run._r.remove(run._r.get_or_add_rPr())
                run._r.insert(0, copy.deepcopy(original_rPr))

    def set_properties(self, title=''):
        """문서 속성: 템플릿에서 물려받은 값 대신 곡 제목과 현재 시각만 남긴다 (작성자 등은 비움)."""
        cp = self.prs.core_properties
        now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
        cp.title = title
        cp.author = ''
        cp.last_modified_by = ''
        cp.revision = 1
        cp.created = now
        cp.modified = now
        # 템플릿의 미리보기 그림은 다른 곡의 것이므로 관계를 끊는다
        package = self.prs.part.package
        for rId, rel in list(package._rels.items()):
            if rel.reltype.endswith('/metadata/thumbnail'):
                package.drop_rel(rId)

    def generate_slides_to_memory(self, lyrics_data, title=''):
        if not self.prs or not self.prs.slides: return None
        self.set_properties(title)

        template_slide = self.prs.slides[0]
        slide_layout = template_slide.slide_layout

        for i in range(len(self.prs.slides) - 1, 0, -1):
            rId = self.prs.slides._sldIdLst[i].rId
            self.prs.part.drop_rel(rId)
            del self.prs.slides._sldIdLst[i]

        for chunk in lyrics_data:
            chunk = clean_lyrics_text(chunk)
            new_slide = self.prs.slides.add_slide(slide_layout)
            for shape in new_slide.shapes:
                 if shape.is_placeholder:
                    sp = shape.element
                    sp.getparent().remove(sp)
            self.duplicate_slide_elements(template_slide, new_slide)
            self.replace_text_in_slide(new_slide, chunk)

        self.delete_template_slide(0)
        
        ppt_buffer = io.BytesIO()
        self.prs.save(ppt_buffer)
        ppt_buffer.seek(0)
        return ppt_buffer

    def delete_template_slide(self, index):
        try:
            rId = self.prs.slides._sldIdLst[index].rId
            self.prs.part.drop_rel(rId)
            del self.prs.slides._sldIdLst[index]
        except: pass

    @staticmethod
    def save_to_file(buffer, output_path):
        """같은 폴더의 임시 파일에 쓴 뒤 바꿔치기해서, 쓰다 실패해도 기존 파일이 깨지지 않게 한다.
        실패하면 임시 파일을 지우고 OSError를 그대로 올린다."""
        folder, name = os.path.split(os.path.abspath(output_path))
        tmp_path = os.path.join(folder, f'.{name}.{os.getpid()}.tmp')
        try:
            with open(tmp_path, 'wb') as f:
                f.write(buffer.getbuffer())
                f.flush()
                os.fsync(f.fileno())
            os.replace(tmp_path, output_path)
        except BaseException:
            try:
                os.remove(tmp_path)
            except OSError:
                pass
            raise
        return True

    @staticmethod
    def split_lyrics(full_lyrics, lines_per_slide=2):
        lines = [l.strip() for l in full_lyrics.split('\n') if l.strip()]
        return ["\n".join(lines[i:i+lines_per_slide]) for i in range(0, len(lines), lines_per_slide)]