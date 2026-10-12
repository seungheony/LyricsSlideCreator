"""PPT에서 가사 한 줄이 글상자 폭을 넘어 두 줄로 나뉘는지(줄바꿈) 미리 계산한다.

템플릿이 바뀌어도 맞도록 값을 코드에 두지 않고, 템플릿 첫 슬라이드의 가사 글상자
(ppt_engine.find_lyrics_shape — PPT 생성과 같은 글상자)에 실제로 적용되는 값을 PowerPoint의 상속 순서대로 읽는다.
- 글상자(폭·안쪽 여백·줄바꿈 여부·세로 쓰기·단·자동 맞춤): 슬라이드 → 레이아웃 개체 틀 → 마스터 개체 틀
- 문단(들여쓰기·글머리 기호)·글자(글꼴·크기·굵게·기울임·자간·커닝·대문자):
  런 → 문단 → 글상자 목록 스타일 → 레이아웃 개체 틀 → 마스터 개체 틀 → 마스터 텍스트 스타일 → 프레젠테이션 기본값
- 테마 글꼴(+mn-lt, +mj-ea 등)은 테마의 글꼴 구성에서 실제 이름을 찾는다
폭은 PowerPoint가 줄을 나누는 방식 그대로 계산한다 (PowerPoint 16.109 Mac에서 줄 수·그려진 폭을 읽어 확인한 규칙):
- 글자마다 글꼴의 설계 폭을 1/8pt(PowerPoint 내부 단위 1/576인치)로 반올림해서 더한다
- 커닝은 글꼴에 예전 방식 'kern' 표가 있을 때만 그 값을 더한다 (요즘 방식 GPOS 커닝은 PowerPoint가 쓰지 않는다)
- 합이 쓸 수 있는 폭을 넘으면 줄을 나눈다. 단, 문장부호 여백 밖 허용(hangingPunct)이면
  줄 끝 문장부호 한 글자(HANGING_PUNCT)는 여백 밖으로 나가도 된다
한글·한자·가나 등은 동아시아 글꼴, 나머지는 라틴 글꼴의 폭을 쓴다.
글꼴이 이 컴퓨터(앱에 들어 있는 글꼴 포함)에 없으면 비슷한 글꼴로 재므로 어림값이 된다(TextBox.exact).
템플릿에 내장된 글꼴은 압축(MicroType Express)되어 있어 쓰지 않는다.
"""
import math
import struct
from dataclasses import dataclass

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from PySide6.QtGui import QFont, QFontDatabase, QFontInfo, QFontMetricsF, QRawFont

from src.utils.ppt_engine import find_lyrics_shape

NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
EMU_PER_PT = 12700
DEFAULT_INSET = 91440             # PowerPoint 기본 좌우 안쪽 여백 (0.1인치)
DEFAULT_SIZE = 1800               # 글자 크기 기본값 (1/100pt)
GRID_PER_PT = 8                   # PowerPoint는 글자 폭을 1/8pt 단위로 반올림한다
# 문장부호 여백 밖 허용일 때 줄 끝에서 여백 밖으로 나갈 수 있는 문장부호 (PowerPoint에서 확인한 글자만)
# 나가지 않는 것: ' " … · ~ - % ♪ 등
HANGING_PUNCT = set('.,?!;:)]}。、”’」') | set('，．？！：；）］｝』】〉》〕')

# 개체 틀 종류 → 물려받는 마스터 개체 틀 종류 (그 외는 body)
MASTER_PH_TYPE = {'title': 'title', 'ctrTitle': 'title', 'dt': 'dt', 'ftr': 'ftr', 'sldNum': 'sldNum', 'hdr': 'hdr'}
# 마스터 개체 틀 종류 → 마스터 텍스트 스타일
MASTER_TEXT_STYLE = {'title': 'p:titleStyle', 'body': 'p:bodyStyle'}   # 그 외(날짜·바닥글 등)는 otherStyle
BULLETS = ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buBlip')
AUTOFITS = ('a:noAutofit', 'a:normAutofit', 'a:spAutoFit')

# 동아시아 글꼴로 그리는 문자 (한글·한자·가나·전각 문자 등)
EAST_ASIAN_RANGES = (
    (0x1100, 0x11FF), (0x2E80, 0x2FDF), (0x3000, 0x33FF), (0x3400, 0x4DBF), (0x4E00, 0x9FFF),
    (0xA960, 0xA97F), (0xAC00, 0xD7FF), (0xF900, 0xFAFF), (0xFE30, 0xFE4F), (0xFF00, 0xFFEF),
    (0x20000, 0x2FA1F),
)


@dataclass
class TextBox:
    """가사 글상자에서 한 줄의 폭을 재는 데 필요한 값"""
    available_pt: float      # 한 줄에 쓸 수 있는 폭 (안쪽 여백·들여쓰기·단을 뺀 값)
    size_pt: float           # 글자 크기 (자동 맞춤으로 줄어든 비율 포함)
    bold: bool
    italic: bool
    latin: str               # 라틴 글꼴 이름
    east_asian: str          # 동아시아 글꼴 이름
    spacing_pt: float        # 자간 (글자마다 더해짐)
    kerning: bool
    caps: bool               # 모두 대문자로 표시
    wraps: bool              # False면 PowerPoint가 줄을 나누지 않는다 (줄바꿈 없음 설정, 세로 쓰기)
    hanging: bool            # 문장부호 여백 밖 허용 (줄 끝 문장부호 한 글자는 폭에서 빠진다)
    exact: bool = True       # 글꼴이 이 컴퓨터에 있어 정확히 잴 수 있는지


def _is_east_asian(ch):
    code = ord(ch)
    return any(lo <= code <= hi for lo, hi in EAST_ASIAN_RANGES)


def _first(elements, attr, default=None):
    """상속 순서대로 놓인 요소들에서 처음으로 지정된 값"""
    for el in elements:
        if el is not None and el.get(attr) is not None:
            return el.get(attr)
    return default


def _ph(sp):
    """개체 틀이면 (종류, 번호), 아니면 None"""
    ph = sp.find('p:nvSpPr/p:nvPr/p:ph', NS)
    if ph is None:
        return None
    return ph.get('type', 'obj'), int(ph.get('idx', '0'))


def _match_ph(tree, ph_type, ph_idx, by_master_type=False):
    """레이아웃·마스터에서 같은 개체 틀 찾기: 번호가 있으면 번호로, 아니면 종류로"""
    candidates = [(sp, _ph(sp)) for sp in tree.iter('{%s}sp' % NS['p'])]
    candidates = [(sp, ph) for sp, ph in candidates if ph]
    if ph_idx and not by_master_type:
        for sp, (_, idx) in candidates:
            if idx == ph_idx:
                return sp
    want = MASTER_PH_TYPE.get(ph_type, 'body') if by_master_type else ph_type
    for sp, (kind, _) in candidates:
        if (MASTER_PH_TYPE.get(kind, 'body') if by_master_type else kind) == want:
            return sp
    if not by_master_type:
        return _match_ph(tree, ph_type, ph_idx, by_master_type=True)
    return None


def _theme_fonts(master):
    """테마 글꼴 구성: {'mj': {'lt': 이름, 'ea': 이름}, 'mn': {...}}"""
    theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    fonts = {}
    for key, tag in (('mj', 'a:majorFont'), ('mn', 'a:minorFont')):
        el = theme.find('.//a:fontScheme/' + tag, NS)
        names = {'lt': '', 'ea': ''}
        if el is not None:
            for script, child in (('lt', 'a:latin'), ('ea', 'a:ea')):
                c = el.find(child, NS)
                names[script] = c.get('typeface', '') if c is not None else ''
            if not names['ea']:
                # 동아시아 글꼴이 비어 있으면 한글 문자에는 언어별 글꼴(Hang)을 쓴다
                hang = el.find("a:font[@script='Hang']", NS)
                names['ea'] = hang.get('typeface', '') if hang is not None else ''
        fonts[key] = names
    return fonts


def _resolve_typeface(name, theme_fonts):
    """'+mn-lt'·'+mj-ea' 같은 테마 글꼴 참조를 실제 글꼴 이름으로"""
    if name and name.startswith('+') and len(name) >= 6:
        return theme_fonts.get(name[1:3], {}).get(name[4:6], '')
    return name or ''


def read_text_box(template_path):
    """템플릿의 가사 글상자 설정을 읽는다. 가사 글상자가 없으면 None"""
    prs = Presentation(template_path)
    if not prs.slides:
        return None
    slide = prs.slides[0]
    shape = find_lyrics_shape(slide)
    if shape is None:
        return None
    layout = slide.slide_layout
    master = layout.slide_master
    sp = shape.element

    # 개체 틀이면 레이아웃·마스터의 같은 개체 틀에서 물려받는다
    layout_sp = master_sp = text_style = None
    ph = _ph(sp)
    if ph:
        layout_sp = _match_ph(layout._element, *ph)
        layout_ph = _ph(layout_sp) if layout_sp is not None else ph
        master_sp = _match_ph(master._element, *layout_ph, by_master_type=True)
        master_type = MASTER_PH_TYPE.get(layout_ph[0], 'body')
        text_style = master._element.find('p:txStyles/' + MASTER_TEXT_STYLE.get(master_type, 'p:otherStyle'), NS)
    shapes = [sp, layout_sp, master_sp]

    # ── 글상자 ──
    body = [s.find('p:txBody/a:bodyPr', NS) if s is not None else None for s in shapes]
    width = _first([s.find('p:spPr/a:xfrm/a:ext', NS) if s is not None else None for s in shapes], 'cx')
    if width is None:
        return None
    insets = int(_first(body, 'lIns', DEFAULT_INSET)) + int(_first(body, 'rIns', DEFAULT_INSET))
    columns = max(1, int(_first(body, 'numCol', 1)))
    column_gap = int(_first(body, 'spcCol', 0))
    column_width = (int(width) - insets - (columns - 1) * column_gap) / columns
    wraps = _first(body, 'wrap', 'square') != 'none' and _first(body, 'vert', 'horz') == 'horz'
    font_scale = 1.0
    for b in body:
        fit = next((b.find(tag, NS) for tag in AUTOFITS if b is not None and b.find(tag, NS) is not None), None)
        if fit is not None:
            # 글자 줄이기(normAutofit)로 이미 줄어든 비율이 있으면 반영한다
            if fit.tag.endswith('normAutofit'):
                font_scale = int(fit.get('fontScale', '100000')) / 100000
            break

    # ── 문단·글자: 상속 순서대로 ──
    paragraph = shape.text_frame.paragraphs[0]._p
    run = shape.text_frame.paragraphs[0].runs[0]._r
    p_pr = paragraph.find('a:pPr', NS)
    level = 'a:lvl%dpPr' % (int(_first([p_pr], 'lvl', 0)) + 1)
    default_style = prs._element.find('p:defaultTextStyle', NS)
    styles = [s.find('p:txBody/a:lstStyle', NS) if s is not None else None for s in shapes] + [text_style, default_style]
    levels = [p_pr] + [s.find(level, NS) if s is not None else None for s in styles]
    if default_style is not None:
        levels.append(default_style.find('a:defPPr', NS))
    chars = [run.find('a:rPr', NS)] + [lv.find('a:defRPr', NS) if lv is not None else None for lv in levels]

    def typeface(child, default):
        for el in chars:
            c = el.find(child, NS) if el is not None else None
            if c is not None and c.get('typeface'):
                return c.get('typeface')
        return default

    has_bullet = False
    for lv in levels:
        bullet = next((lv.find(tag, NS) for tag in BULLETS if lv is not None and lv.find(tag, NS) is not None), None)
        if bullet is not None:
            has_bullet = not bullet.tag.endswith('buNone')
            break
    margin = int(_first(levels, 'marL', 0))
    indent = int(_first(levels, 'indent', 0))
    # 첫 줄 글자가 시작하는 곳: 글머리 기호가 있으면 기호 뒤(왼쪽 여백), 없으면 첫 줄 들여쓰기만큼
    start = margin + max(indent, 0) if has_bullet else margin + indent
    available = column_width - start - int(_first(levels, 'marR', 0))

    theme_fonts = _theme_fonts(master)
    latin = _resolve_typeface(typeface('a:latin', '+mn-lt'), theme_fonts)
    east_asian = _resolve_typeface(typeface('a:ea', '+mn-ea'), theme_fonts) or latin
    size = int(_first(chars, 'sz', DEFAULT_SIZE))
    kern = _first(chars, 'kern')
    return TextBox(
        available_pt=available / EMU_PER_PT,
        size_pt=size / 100 * font_scale,
        bold=_first(chars, 'b', '0') in ('1', 'true'),
        italic=_first(chars, 'i', '0') in ('1', 'true'),
        latin=latin,
        east_asian=east_asian,
        spacing_pt=int(_first(chars, 'spc', 0)) / 100,
        # kern: 이 크기(1/100pt) 이상에서 커닝. 없으면 모든 크기에서 커닝
        kerning=kern is None or size >= int(kern),
        caps=_first(chars, 'cap', 'none') == 'all',
        hanging=_first(levels, 'hangingPunct', '0') in ('1', 'true'),
        wraps=wraps and available > 0,
    )


def _font_available(family):
    """이 컴퓨터에서 그 이름의 글꼴을 쓸 수 있는지 (한글 이름도 Qt가 찾아준다)"""
    if not family:
        return False
    if QFontDatabase.hasFamily(family):
        return True
    # 없는 이름은 기본 글꼴로 대신 잡히므로, 대신 잡힌 글꼴과 다르면 있는 것
    substitute = QFontInfo(QFont('LSC-NoSuchFont-7f3a')).family()
    return QFontInfo(QFont(family)).family() != substitute


def _legacy_kerning(raw):
    """예전 방식 커닝 표('kern', Microsoft 형식 0)의 글자 쌍 값 {(왼쪽 글리프, 오른쪽 글리프): 글꼴 단위}.
    PowerPoint는 이 표만 쓰고 GPOS 커닝은 쓰지 않는다 (Arial은 커닝됨, GPOS만 있는 Pretendard는 안 됨)"""
    data = bytes(raw.fontTable('kern'))
    pairs = {}
    if len(data) < 4 or struct.unpack_from('>H', data, 0)[0] != 0:
        return pairs
    offset = 4
    for _ in range(struct.unpack_from('>H', data, 2)[0]):
        if offset + 14 > len(data):
            break
        length, coverage = struct.unpack_from('>HH', data, offset + 2)
        if coverage >> 8 != 0:
            offset += length
            continue
        count = struct.unpack_from('>H', data, offset + 6)[0]
        # 가로 방향만 (최소값·교차 방향 표는 아님). 길이 칸은 넘칠 수 있어 쌍 개수로 끝을 구한다
        if coverage & 0x1 and not coverage & 0x6:
            for i in range(count):
                left, right, value = struct.unpack_from('>HHh', data, offset + 14 + i * 6)
                pairs[(left, right)] = value if coverage & 0x8 else pairs.get((left, right), 0) + value
        offset += 14 + count * 6
    return pairs


class _Face:
    """한 글꼴의 글자별 폭(1/8pt로 반올림)과 커닝 값"""

    def __init__(self, family, box):
        font = QFont(family)
        font.setPixelSize(100)
        font.setBold(box.bold)
        font.setItalic(box.italic)
        font.setHintingPreference(QFont.HintingPreference.PreferNoHinting)
        self.raw = QRawFont.fromFont(font)
        units = self.raw.unitsPerEm()
        self.raw.setPixelSize(units)          # 이 크기에서 폭 = 글꼴 설계 단위
        self.pt_per_unit = box.size_pt / units
        self.pairs = _legacy_kerning(self.raw) if box.kerning else {}
        # 글꼴에 없는 글자는 Qt가 대신 고른 글꼴로 잰다
        font.setPixelSize(max(1, round(box.size_pt * 32)))
        self.fallback = QFontMetricsF(font)
        self.fallback_scale = box.size_pt / font.pixelSize()
        self.glyphs = {}

    def glyph(self, ch):
        """(글리프 번호, 1/8pt로 반올림한 폭)"""
        if ch not in self.glyphs:
            ids = self.raw.glyphIndexesForString(ch)
            gid = ids[0] if ids else 0
            if gid:
                units = self.raw.advancesForGlyphIndexes([gid], QRawFont.LayoutFlag.UseDesignMetrics)[0].x()
                width = units * self.pt_per_unit
            else:
                width = self.fallback.horizontalAdvance(ch) * self.fallback_scale
            self.glyphs[ch] = (gid, math.floor(width * GRID_PER_PT + 0.5) / GRID_PER_PT)
        return self.glyphs[ch]


class LineFit:
    """가사 줄마다 PPT 글상자 폭 대비 글자 폭 비율을 잰다 (1을 넘으면 PPT에서 줄바꿈된다).
    QFont를 쓰므로 QGuiApplication이 만들어진 뒤 GUI 스레드에서 쓴다."""

    def __init__(self, template_path):
        self.box = read_text_box(template_path)
        self._faces = {}
        self._cache = {}
        if self.box:
            self.box.exact = _font_available(self.box.latin) and _font_available(self.box.east_asian)

    def _face(self, family):
        if family not in self._faces:
            self._faces[family] = _Face(family, self.box)
        return self._faces[family]

    def width_pt(self, text):
        """PowerPoint가 이 줄을 그리는 폭 (pt): 글자별 반올림 폭 + 커닝 + 자간"""
        box = self.box
        if box.caps:
            text = text.upper()
        width = 0.0
        prev = None
        for ch in text:
            family = box.east_asian if _is_east_asian(ch) else box.latin
            face = self._face(family)
            gid, advance = face.glyph(ch)
            width += advance
            if prev and prev[0] == family and face.pairs:
                width += face.pairs.get((prev[1], gid), 0) * face.pt_per_unit
            prev = (family, gid)
        return width + box.spacing_pt * len(text)

    def ratio(self, text):
        """글상자 폭 대비 비율. 1을 넘으면 PPT에서 두 줄 이상으로 나뉜다. 계산할 수 없으면 0"""
        if not self.box or not self.box.wraps or not text:
            return 0.0
        if text not in self._cache:
            width = self.width_pt(text)
            if width > self.box.available_pt and self.box.hanging and text[-1] in HANGING_PUNCT:
                # 줄 끝 문장부호 한 글자는 여백 밖으로 나가도 된다
                width = self.width_pt(text[:-1])
            # 반올림하지 않는다: 1/8pt만 넘어도(705.625/705.6 = 1.00004) 줄바꿈이다
            self._cache[text] = width / self.box.available_pt
        return self._cache[text]
