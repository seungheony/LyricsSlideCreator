"""다른 PPT(앞·뒤에 붙일 PPT)의 슬라이드를 가사 PPT에 끼워 넣는다.

PowerPoint의 '원본 서식 유지' 붙여넣기처럼, 슬라이드와 함께 그 슬라이드가 쓰는 레이아웃·마스터·테마·
그림·영상·차트 등 연결된 파트를 모두 복사한다(파일 하나 안에서는 한 번씩만).
- 슬라이드 노트는 붙이는 쪽 PPT의 노트 마스터를 쓴다(노트 마스터는 PPT에 하나뿐)
- 메모(댓글)는 작성자 목록이 PPT 단위라 가져오지 않는다
- 구역(섹션)·내장 글꼴처럼 PPT 전체에 붙은 설정은 가져오지 않는다
- 슬라이드 크기는 PPT에 하나뿐이라 처음 붙이는 PPT 크기를 따르고 가사 슬라이드를 거기에 맞춘다(fit_to_size).
  붙이는 PPT끼리 크기가 다르면 그대로 붙인다(화면에서 미리 알린다)
"""
import posixpath
import re
import zipfile

from lxml import etree

from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.constants import RELATIONSHIP_TARGET_MODE as RTM
from pptx.opc.package import PartFactory, _Relationship
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn

# 가져오지 않는 관계: 메모(예전·새 방식)와 그 작성자
_SKIP_REL = re.compile(r'/(comments|commentAuthors|authors)$')
# 마스터·레이아웃 id는 2^31부터 (ECMA-376), 마스터와 레이아웃이 같은 범위를 나눠 쓴다
_MIN_MASTER_ID = 2147483648


def deck_info(path):
    """붙일 PPT 미리 보기: {'slides': 장 수, 'width', 'height'(EMU)}.
    그림·영상까지 읽지 않도록 presentation.xml만 본다(설교 PPT는 클 수 있다). 열 수 없으면 ValueError"""
    try:
        with zipfile.ZipFile(path) as z:
            rels = etree.fromstring(z.read('_rels/.rels'))
            main = next(r.get('Target') for r in rels
                        if r.get('Type', '').endswith('/officeDocument'))
            root = etree.fromstring(z.read(main.lstrip('/')))
    except (OSError, KeyError, StopIteration, zipfile.BadZipFile, etree.XMLSyntaxError) as e:
        raise ValueError(str(e)) from e
    if root.tag != qn('p:presentation'):
        raise ValueError('PowerPoint 파일이 아님')
    size = root.find(qn('p:sldSz'))
    if size is None:
        raise ValueError('슬라이드 크기 정보가 없음')
    return {'slides': len(root.findall(f"{qn('p:sldIdLst')}/{qn('p:sldId')}")),
            'width': int(size.get('cx')), 'height': int(size.get('cy'))}


class _Copier:
    """한 원본 PPT의 파트를 대상 PPT로 복사한다. 같은 파트는 한 번만 복사한다(memo)"""

    def __init__(self, dest):
        self.dest = dest
        self.package = dest.part.package
        self.used = {str(p.partname).lower() for p in self.package.iter_parts()}
        self.memo = {}

    def _partname(self, src_name):
        folder, base = posixpath.split(str(src_name))
        stem, ext = posixpath.splitext(base)
        stem = stem.rstrip('0123456789') or 'part'
        n = 1
        while f'{folder}/{stem}{n}{ext}'.lower() in self.used:
            n += 1
        name = f'{folder}/{stem}{n}{ext}'
        self.used.add(name.lower())
        return PackURI(name)

    def copy(self, part):
        if part in self.memo:
            return self.memo[part]
        new = PartFactory(self._partname(part.partname), part.content_type, self.package, part.blob)
        self.memo[part] = new          # 순환 관계(레이아웃 ↔ 마스터, 노트 ↔ 슬라이드)를 위해 먼저 등록
        rels = new.rels
        # XML이 rId로 가리키므로 rId를 그대로 둔다
        for rId, rel in part.rels.items():
            if _SKIP_REL.search(rel.reltype):
                continue
            if rel.is_external:
                target, mode = rel.target_ref, RTM.EXTERNAL
            else:
                src = rel.target_part
                if src.content_type == CT.PML_NOTES_MASTER:
                    target = self.dest.part.notes_master_part
                elif src is part.package.main_document_part:
                    continue
                else:
                    target = self.copy(src)
                mode = RTM.INTERNAL
            rels._rels[rId] = _Relationship(rels._base_uri, rId, rel.reltype, mode, target)
        return new


def _register_masters(dest, parts):
    """복사한 마스터를 presentation.xml 목록에 넣고, 마스터·레이아웃 id를 겹치지 않게 다시 매긴다"""
    masters = [p for p in parts if p.content_type == CT.PML_SLIDE_MASTER]
    if not masters:
        return
    lst = dest.part._element.get_or_add_sldMasterIdLst()
    used = [_MIN_MASTER_ID - 1]
    for entry in lst.sldMasterId_lst:
        used.append(int(entry.get('id')))
        master = dest.part.related_part(entry.rId)
        used += [int(e.get('id')) for e in master._element.iter(qn('p:sldLayoutId'))]
    next_id = max(used) + 1
    for master in masters:
        entry = lst._add_sldMasterId()
        entry.set('id', str(next_id))
        entry.set(qn('r:id'), dest.part.relate_to(master, RT.SLIDE_MASTER))
        next_id += 1
        for layout_id in master._element.iter(qn('p:sldLayoutId')):
            layout_id.set('id', str(next_id))
            next_id += 1


# 크기를 맞출 때 바꾸는 값 (모두 EMU 또는 1/100pt — 비율로 늘리고 줄인다)
_SCALE_ATTRS = {
    'off': ('x', 'y'), 'chOff': ('x', 'y'), 'ext': ('cx', 'cy'), 'chExt': ('cx', 'cy'),
    'rPr': ('sz', 'kern', 'spc'), 'defRPr': ('sz', 'kern', 'spc'), 'endParaRPr': ('sz', 'kern', 'spc'),
    'spcPts': ('val',), 'buSzPts': ('val',), 'tab': ('pos',), 'ln': ('w',),
    'bodyPr': ('lIns', 'tIns', 'rIns', 'bIns'),
}
for _lvl in ['defPPr', 'pPr'] + [f'lvl{i}pPr' for i in range(1, 10)]:
    _SCALE_ATTRS[_lvl] = ('marL', 'indent', 'defTabSz')
_A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
_INSETS = {'lIns': 91440, 'rIns': 91440, 'tIns': 45720, 'bIns': 45720}   # 글상자 안쪽 여백 기본값


def _scale_xml(root, k, dx, dy, is_master):
    """도형 위치·크기·글자 크기·여백을 k배로 하고, 맨 위 도형은 (dx, dy)만큼 옮긴다(가운데 맞춤)"""
    sp_tree = root.find('.//' + qn('p:spTree'))
    top = set()
    if sp_tree is not None:
        for shape in sp_tree:
            if shape.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')):
                continue
            xfrm = next(shape.iter(qn('a:xfrm'), qn('p:xfrm')), None)
            if xfrm is not None:
                top.update(xfrm.findall(qn('a:off')))
    for el in root.iter():
        if not isinstance(el.tag, str) or not el.tag.startswith(_A):
            continue
        name = el.tag[len(_A):]
        if name == 'bodyPr':
            # 상속받지 않는 글상자(마스터, 개체 틀이 아닌 도형)는 기본 여백도 늘려야 하므로 적어 둔다
            sp = el.getparent().getparent()
            is_ph = sp is not None and sp.find('.//' + qn('p:ph')) is not None
            if is_master or not is_ph:
                for attr, value in _INSETS.items():
                    if el.get(attr) is None:
                        el.set(attr, str(value))
        for attr in _SCALE_ATTRS.get(name, ()):
            value = el.get(attr)
            if value is None or not value.lstrip('-').isdigit():
                continue
            v = round(int(value) * k)
            if el in top:
                v += dx if attr == 'x' else dy
            if attr == 'sz':
                v = min(max(v, 100), 400000)
            el.set(attr, str(v))


def fit_to_size(prs, width, height):
    """가사 PPT를 (width, height) 크기로 바꾼다. 비율이 다르면 작은 쪽에 맞춰 가운데에 둔다(배경은 슬라이드 전체).
    가사 PPT의 마스터·레이아웃·슬라이드를 모두 같은 비율로 바꾸므로 줄바꿈 위치도 그대로다."""
    w0, h0 = prs.slide_width, prs.slide_height
    if (w0, h0) == (width, height):
        return
    k = min(width / w0, height / h0)
    dx, dy = round((width - w0 * k) / 2), round((height - h0 * k) / 2)
    parts = []
    for master in prs.slide_masters:
        parts.append((master.part, True))
        parts += [(layout.part, False) for layout in master.slide_layouts]
    parts += [(slide.part, False) for slide in prs.slides]
    for part, is_master in parts:
        _scale_xml(part._element, k, dx, dy, is_master)
    prs.slide_width, prs.slide_height = width, height


def insert_deck(dest, src, index):
    """src(Presentation)의 슬라이드를 모두 dest의 index 위치부터 차례로 넣는다. 반환: 넣은 장 수"""
    copier = _Copier(dest)
    new_slides = [copier.copy(slide.part) for slide in src.slides]
    _register_masters(dest, copier.memo.values())
    sld_lst = dest.slides._sldIdLst
    for i, part in enumerate(new_slides):
        sld_id = sld_lst.add_sldId(dest.part.relate_to(part, RT.SLIDE))
        sld_lst.remove(sld_id)
        sld_lst.insert(index + i, sld_id)
    return len(new_slides)


def attach_decks(dest, before=(), after=()):
    """가사 슬라이드 앞에 before, 뒤에 after의 PPT들(파일 경로)을 순서대로 붙인다. 반환: 붙인 장 수.
    슬라이드 크기는 처음 붙이는 PPT를 따르고 가사 슬라이드를 거기에 맞춘다(fit_to_size)."""
    before = [Presentation(p) for p in before]
    after = [Presentation(p) for p in after]
    decks = before + after
    if not decks:
        return 0
    first = decks[0]
    fit_to_size(dest, first.slide_width, first.slide_height)
    # 크기 종류(screen16x9 등)도 처음 PPT를 따른다 (없으면 사용자 지정 크기)
    size_type = first.part._element.sldSz.get('type')
    dest_size = dest.part._element.sldSz
    if size_type:
        dest_size.set('type', size_type)
    elif 'type' in dest_size.attrib:
        del dest_size.attrib['type']
    pos = 0
    for src in before:
        pos += insert_deck(dest, src, pos)
    for src in after:
        insert_deck(dest, src, len(dest.slides._sldIdLst))
    dest.part.rename_slide_parts([s.rId for s in dest.slides._sldIdLst])
    return sum(len(d.slides) for d in decks)
