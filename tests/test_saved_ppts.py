"""보관한 PPT(saved_ppts)와 앞·뒤 목록에서 보관본을 쓰는 규칙(backend)을 시험한다."""
import json
import os
import shutil
import tempfile
import unittest

from pptx import Presentation

from src.qml_ui.backend import attach_path, clean_attach_lists, describe_attachments
from src.utils import app_state, saved_ppts


class SavedPptsTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix='lsc-saved-')
        self.data = os.path.join(self.dir, 'data')
        self.old_env = os.environ.get('LSC_DATA_DIR')
        os.environ['LSC_DATA_DIR'] = self.data
        self.src = os.path.join(self.dir, '예배 준비중.pptx')
        self.make(self.src, 2)

    def tearDown(self):
        os.environ['LSC_DATA_DIR'] = self.old_env
        shutil.rmtree(self.dir, ignore_errors=True)

    @staticmethod
    def make(path, slides):
        prs = Presentation()
        for _ in range(slides):
            prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(path)

    def test_add_copies_and_survives_original(self):
        items, item, replaced = saved_ppts.add([], self.src, '2026-10-16T10:00:00')
        self.assertFalse(replaced)
        self.assertEqual(item['name'], '예배 준비중.pptx')
        self.assertTrue(saved_ppts.save(items))
        os.remove(self.src)
        loaded = saved_ppts.load()
        self.assertEqual(loaded, items)
        info = describe_attachments({'before': [{'saved': item['id']}]}, loaded)
        self.assertEqual(info['problems'], 0)
        self.assertEqual(info['before'][0]['slides'], 2)
        self.assertTrue(info['before'][0]['saved'])
        self.assertEqual(info['before'][0]['name'], '예배 준비중.pptx')
        self.assertEqual(attach_path({'saved': item['id']}), saved_ppts.file_path(item['id']))

    def test_same_name_replaces(self):
        items, first, _ = saved_ppts.add([], self.src, 'a')
        self.make(self.src, 3)                     # 원본을 고친 뒤 다시 보관
        items, second, replaced = saved_ppts.add(items, self.src, 'b')
        self.assertTrue(replaced)
        self.assertEqual(second['id'], first['id'])
        self.assertEqual(len(items), 1)
        self.assertEqual(len(Presentation(saved_ppts.file_path(first['id'])).slides), 3)

    def test_remove_deletes_file(self):
        items, item, _ = saved_ppts.add([], self.src, 'a')
        items = saved_ppts.remove(items, item['id'])
        self.assertEqual(items, [])
        self.assertFalse(os.path.exists(saved_ppts.file_path(item['id'])))
        self.assertTrue(app_state.data_dir().startswith(self.data))
        info = describe_attachments({'after': [{'saved': item['id']}]}, items)
        self.assertEqual(info['problems'], 1)
        self.assertTrue(info['after'][0]['missing'])

    def test_source_and_stale(self):
        items, item, _ = saved_ppts.add([], self.src, 'a')
        self.assertEqual(item['source'], self.src)
        lists = {'before': [{'saved': item['id']}], 'after': [self.src]}
        now = os.path.getmtime(self.src)
        info = describe_attachments(lists, items, now=now + 3 * 86400)
        self.assertTrue(info['before'][0]['hasSource'])
        self.assertNotIn('stale', info['after'][0])
        info = describe_attachments(lists, items, now=now + 7 * 86400)
        self.assertIn('월', info['after'][0]['stale'])          # 보관하지 않은 파일만 알린다
        self.assertNotIn('stale', info['before'][0])
        self.assertNotIn('modified', info['after'][0])
        os.remove(self.src)
        self.assertFalse(describe_attachments(lists, items)['before'][0]['hasSource'])

    def test_rename(self):
        items, item, _ = saved_ppts.add([], self.src, 'a')
        renamed = saved_ppts.rename(items, item['id'], '  설교   양식 ')
        self.assertEqual(renamed[0]['name'], '설교 양식')
        self.assertIsNone(saved_ppts.rename(items, item['id'], '   '))
        self.assertIsNone(saved_ppts.rename(items, 'nope', 'x'))

    def test_backend_library_only(self):
        from src.qml_ui.backend import AlbumArtProvider, Backend
        backend = Backend(AlbumArtProvider())
        res = json.loads(backend.savePptToLibrary(self.src))
        self.assertIn('item', res)
        self.assertEqual(json.loads(backend.attachments())['before'], [])      # 목록에는 넣지 않는다
        self.assertEqual([i['name'] for i in json.loads(backend.savedPpts())], ['예배 준비중.pptx'])
        self.assertTrue(backend.renameSavedPpt(res['item']['id'], '예배 준비'))
        self.assertEqual(json.loads(backend.savedPpts())[0]['name'], '예배 준비')
        bad = json.loads(backend.savePptToLibrary(os.path.join(self.dir, 'x.txt')))
        self.assertIn('error', bad)

    def test_clean_lists(self):
        self.assertEqual(clean_attach_lists({'before': ['/a.pptx', {'saved': 'x', 'junk': 1}, 3, {'saved': 4}],
                                             'after': None}),
                         {'before': ['/a.pptx', {'saved': 'x'}], 'after': []})
        self.assertEqual(clean_attach_lists(None), {'before': [], 'after': []})


if __name__ == '__main__':
    unittest.main()
