import io
import os
import unittest

from pptx import Presentation

from src.utils.ppt_engine import PPTBuilder, clean_lyrics_text

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, 'template.pptx')


def lyrics_shape(slide):
    """생성된 슬라이드의 가사 글상자 (가사가 바뀌었으므로 '가사'를 찾는 find_lyrics_shape 대신 첫 글상자)"""
    return next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text)


def slide_texts(prs):
    return [lyrics_shape(s).text_frame.text for s in prs.slides]


class PPTBuilderTest(unittest.TestCase):
    def build(self, chunks, title='시험 곡'):
        buf = PPTBuilder(TEMPLATE).generate_slides_to_memory(chunks, title=title)
        return Presentation(io.BytesIO(buf.getvalue()))

    def test_slide_count_and_text(self):
        prs = self.build(['첫째 줄\n둘째 줄', '셋째 줄', '넷째\n다섯째'])
        self.assertEqual(len(prs.slides), 3)
        self.assertEqual(slide_texts(prs), ['첫째 줄\n둘째 줄', '셋째 줄', '넷째\n다섯째'])

    def test_each_line_is_own_paragraph(self):
        prs = self.build(['가\n나'])
        shape = lyrics_shape(prs.slides[0])
        self.assertEqual([p.text for p in shape.text_frame.paragraphs], ['가', '나'])

    def test_document_properties(self):
        prs = self.build(['가'], title='곡 제목')
        cp = prs.core_properties
        self.assertEqual(cp.title, '곡 제목')
        self.assertEqual(cp.author, '')
        self.assertEqual(cp.last_modified_by, '')
        self.assertIsNotNone(cp.created)
        self.assertEqual(cp.created, cp.modified)

    def test_xml_special_chars_and_emoji(self):
        prs = self.build(['a < b & c > d\n😀 ♪'])
        self.assertEqual(slide_texts(prs), ['a < b & c > d\n😀 ♪'])

    def test_control_chars_cleaned(self):
        self.assertEqual(clean_lyrics_text('가\x00나​다\t라\r\n마 바'), '가나다 라\n마\n바')
        prs = self.build(['가\x0b나'])
        self.assertEqual(slide_texts(prs), ['가\n나'])

    def test_save_to_file_is_atomic(self):
        import tempfile
        buf = PPTBuilder(TEMPLATE).generate_slides_to_memory(['가'])
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, 'out.pptx')
            PPTBuilder.save_to_file(buf, path)
            self.assertEqual(os.listdir(d), ['out.pptx'])
            self.assertEqual(len(Presentation(path).slides), 1)


if __name__ == '__main__':
    unittest.main()
