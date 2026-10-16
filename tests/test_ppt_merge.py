"""앞·뒤에 붙일 PPT(ppt_merge)와 그 목록 정보(backend.describe_attachments)를 시험한다.
붙일 PPT는 python-pptx 기본 템플릿(가사 템플릿과 다른 마스터)으로 그림·차트·노트·링크·숨긴 슬라이드를 넣어 만든다."""
import io
import os
import shutil
import struct
import tempfile
import unittest
import zipfile
import zlib

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

from src.qml_ui.backend import describe_attachments, inspect_attachment
from src.utils.ppt_engine import PPTBuilder
from src.utils.ppt_merge import deck_info

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, 'template.pptx')
WIDE = (Inches(13.333), Inches(7.5))      # PowerPoint 기본 16:9 (가사 템플릿은 10 × 5.625인치)


def png_bytes():
    """2×2 빨간 PNG"""
    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data))
    raw = b''.join(b'\x00' + b'\xff\x00\x00' * 2 for _ in range(2))
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 2, 2, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


def texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text]


class MergeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.mkdtemp(prefix='lsc-merge-')
        pic = os.path.join(cls.dir, 'pic.png')
        with open(pic, 'wb') as f:
            f.write(png_bytes())

        front = Presentation()
        front.slide_width, front.slide_height = WIDE
        s1 = front.slides.add_slide(front.slide_layouts[0])
        s1.shapes.title.text = '예배 준비중'
        s1.notes_slide.notes_text_frame.text = '앞 노트'
        s2 = front.slides.add_slide(front.slide_layouts[5])
        s2.shapes.title.text = '광고'
        s2.shapes.add_picture(pic, Inches(1), Inches(1), Inches(2))
        link = s2.shapes.add_textbox(Inches(4), Inches(1), Inches(3), Inches(1))
        link.text_frame.text = '처음으로'
        link.click_action.target_slide = s1
        data = CategoryChartData()
        data.categories = ['1부', '2부']
        data.add_series('출석', (120, 80))
        s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(3), Inches(4), Inches(3), data)
        cls.front = os.path.join(cls.dir, 'front.pptx')
        front.save(cls.front)

        back = Presentation()
        back.slide_width, back.slide_height = WIDE
        b1 = back.slides.add_slide(back.slide_layouts[1])
        b1.shapes.title.text = '설교'
        b1.notes_slide.notes_text_frame.text = '설교 노트'
        hidden = back.slides.add_slide(back.slide_layouts[6])
        hidden._element.set('show', '0')
        cls.back = os.path.join(cls.dir, 'back.pptx')
        back.save(cls.back)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, ignore_errors=True)

    def build(self, before=(), after=(), chunks=('첫 줄\n둘째 줄',)):
        songs = [{'title': '곡', 'artist': '', 'chunks': list(chunks)}]
        buf = PPTBuilder(TEMPLATE).generate_setlist_to_memory(songs, '제목', before, after)
        return buf, Presentation(io.BytesIO(buf.getvalue()))

    def test_order_and_size(self):
        _, prs = self.build([self.front], [self.back])
        got = [texts(s) for s in prs.slides]
        self.assertEqual(got, [['예배 준비중'], ['광고', '처음으로'], [], ['첫 줄\n둘째 줄'], [], ['설교'], []])
        # 크기는 처음 붙이는 PPT를 따른다
        self.assertEqual((prs.slide_width, prs.slide_height), WIDE)

    def test_without_attachments_unchanged(self):
        _, prs = self.build()
        self.assertEqual(len(prs.slides), 3)
        self.assertEqual((prs.slide_width, prs.slide_height), (Emu(9144000), Emu(5143500)))

    def test_parts_follow(self):
        buf, prs = self.build([self.front], [self.back, self.front])
        slides = list(prs.slides)
        self.assertEqual(len(slides), 3 + 2 + 2 + 2)
        # 노트·숨김 상태
        self.assertEqual(slides[0].notes_slide.notes_text_frame.text, '앞 노트')
        self.assertEqual(slides[5].notes_slide.notes_text_frame.text, '설교 노트')
        self.assertEqual(slides[6]._element.get('show'), '0')
        # 링크는 같은 파일에서 복사한 슬라이드를 가리킨다 (같은 PPT를 두 번 붙여도 각자)
        for title_idx, link_idx in ((0, 1), (7, 8)):
            link = next(sh for sh in slides[link_idx].shapes if sh.has_text_frame and sh.text_frame.text == '처음으로')
            self.assertIs(link.click_action.target_slide.part, slides[title_idx].part)
        # 그림·차트(내장 엑셀)
        names = zipfile.ZipFile(io.BytesIO(buf.getvalue())).namelist()
        self.assertEqual(len([n for n in names if n.startswith('ppt/media/')]), 2)
        self.assertEqual(len([n for n in names if n.startswith('ppt/embeddings/')]), 2)
        # 슬라이드 파일 이름은 순서대로
        self.assertEqual([s.part.partname for s in slides], [f'/ppt/slides/slide{i}.xml' for i in range(1, 10)])

    def test_master_ids_unique(self):
        _, prs = self.build([self.front], [self.back])
        pres = prs.part._element
        ids = [int(e.get('id')) for e in pres.iter(qn('p:sldMasterId'))]
        self.assertEqual(len(prs.slide_masters), 3)
        for master in prs.slide_masters:
            ids += [int(e.get('id')) for e in master._element.iter(qn('p:sldLayoutId'))]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue(all(i >= 2147483648 for i in ids))
        # 붙인 슬라이드는 자기 마스터의 레이아웃을 쓴다
        self.assertEqual(prs.slides[0].slide_layout.name, 'Title Slide')
        self.assertEqual(prs.slides[0].slide_layout.slide_master, prs.slides[1].slide_layout.slide_master)
        self.assertNotEqual(prs.slides[0].slide_layout.slide_master, prs.slides[3].slide_layout.slide_master)

    def test_lyrics_scaled_to_wide(self):
        _, plain = self.build()
        _, wide = self.build([self.front])
        k = WIDE[0] / plain.slide_width
        a = next(sh for sh in plain.slides[1].shapes if sh.has_text_frame)
        b = next(sh for sh in wide.slides[3].shapes if sh.has_text_frame)
        self.assertAlmostEqual(b.width / a.width, k, places=4)
        self.assertAlmostEqual(b.top / a.top, k, places=3)
        self.assertEqual(b.text_frame.paragraphs[0].runs[0].font.size.pt,
                         round(a.text_frame.paragraphs[0].runs[0].font.size.pt * k))

    def test_lyrics_centered_in_4_3(self):
        square = Presentation()        # python-pptx 기본은 4:3 (10 × 7.5인치)
        square.slides.add_slide(square.slide_layouts[6])
        path = os.path.join(self.dir, 'square.pptx')
        square.save(path)
        _, plain = self.build()
        _, prs = self.build([path])
        self.assertEqual((prs.slide_width, prs.slide_height), (square.slide_width, square.slide_height))
        a = next(sh for sh in plain.slides[1].shapes if sh.has_text_frame)
        b = next(sh for sh in prs.slides[2].shapes if sh.has_text_frame)
        # 폭은 같고(10인치) 위아래 남는 자리의 가운데로
        self.assertEqual(b.width, a.width)
        self.assertEqual(b.top - a.top, (square.slide_height - plain.slide_height) // 2)

    def test_deck_info(self):
        self.assertEqual(deck_info(self.front), {'slides': 2, 'width': WIDE[0], 'height': WIDE[1]})
        with self.assertRaises(ValueError):
            deck_info(os.path.join(ROOT, 'README.md'))
        with self.assertRaises(ValueError):
            deck_info(os.path.join(self.dir, 'nope.pptx'))

    def test_describe_attachments(self):
        old = os.path.join(self.dir, 'old.ppt')
        with open(old, 'wb') as f:
            f.write(b'x')
        info = describe_attachments({'before': [self.front, TEMPLATE], 'after': [old, os.path.join(self.dir, 'gone.pptx')]})
        self.assertEqual(info['slides'], 3)
        self.assertEqual(info['problems'], 2)
        self.assertNotIn('sizeMismatch', info['before'][0])
        self.assertTrue(info['before'][1]['sizeMismatch'])
        self.assertIn('.ppt', info['after'][0]['error'])
        self.assertTrue(info['after'][1]['missing'])
        self.assertEqual(inspect_attachment(self.front)['name'], 'front.pptx')


if __name__ == '__main__':
    unittest.main()
