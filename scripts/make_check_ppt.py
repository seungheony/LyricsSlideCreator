"""Windows PowerPoint 실기 확인용 시험 PPT를 만든다 (docs/WINDOWS_CHECK.md 1단계).

앱과 같은 PPTBuilder·template.pptx로 슬라이드를 만들고, 앱의 LineFit이 예측한 줄바꿈과
PowerPoint가 실제로 그리는 줄바꿈을 사람이 눈으로 비교할 수 있게 한다.
- 경계 줄 슬라이드: 한 슬라이드에 한 줄. 가사 글상자에는 줄만 넣고(앞에 글자를 붙이면 폭이 달라진다),
  예측 비율과 '줄바꿈 예상/안 함'은 슬라이드 아래쪽 별도 글상자와 슬라이드 노트에 적는다
- 3·4·5줄 슬라이드: 글상자 아래로 넘치거나 첫 줄이 잘리는지 확인
- 맨 앞 요약 슬라이드: 전체 예측 표

사용법 (프로젝트 루트에서): python scripts/make_check_ppt.py [-o 출력.pptx]
"""
import argparse
import os
import sys

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402
from PySide6.QtGui import QFontDatabase, QGuiApplication  # noqa: E402

from src.utils.helpers import resource_path  # noqa: E402
from src.utils.line_fit import LineFit  # noqa: E402
from src.utils.ppt_engine import PPTBuilder  # noqa: E402

BAND = (0.97, 1.03)
LABEL_FONT = '맑은 고딕'

# 늘려 가며 경계를 찾을 글자 흐름
STREAMS = {
    '한글': '가을 하늘 아래 우리는 함께 걸었네 바람이 불어와도 따뜻한 그 손을 놓지 않고서 ',
    '영문': 'Amazing grace how sweet the sound that saved a wretch like me I once was lost ',
    '한영 섞임': '주님은 my shepherd 나는 부족함이 없네 Hallelujah 찬양해 영원토록 ',
    '숫자·괄호': '1절 (주님께 감사해) 2절 [영광 돌려요] 3절 "사랑해요" 4절 아멘 ',
}
# 줄 끝 문장부호: 여백 밖 허용(hangingPunct)에서 줄바꿈 여부가 갈리는 글자들
PUNCTS = ['.', ',', '?', '!', ';', ')', '”', '…', '~', '-']
PUNCT_BASES = ['주님의 사랑은 끝이 없어요 우리 모두 노래해요 하늘 높이 찬양해요 영원히',
               'Amazing grace how sweet the sound that saved a wretch like me I once was lost',
               '주님은 my shepherd 나는 부족함이 없네 Hallelujah 찬양해 영원토록 우리 함께']
PER_BOUNDARY = 2       # 경계 한쪽(1 이하·1 초과)마다 몇 개씩 고를지


def boundary_lines(fit, stream):
    """글자 흐름을 한 글자씩 늘리며 비율이 1 바로 아래·위인 줄을 고른다."""
    below, above = [], []
    for n in range(1, len(stream) + 1):
        text = stream[:n].rstrip()
        r = fit.ratio(text)
        if BAND[0] <= r <= 1:
            below.append((r, text))
        elif 1 < r <= BAND[1]:
            above.append((r, text))
    below = sorted(set(below))[-PER_BOUNDARY:]
    above = sorted(set(above))[:PER_BOUNDARY]
    return [t for _, t in below + above]


def raw_ratio(fit, text):
    """문장부호 여백 밖 허용을 적용하지 않은 비율 (줄 전체 글자 폭 기준)"""
    return fit.width_pt(text) / fit.box.available_pt


def punct_lines(fit, punct):
    """끝에 문장부호를 붙인 줄 중 경계(0.97~1.03)인 것: 1 이하 쪽 끝, 문장부호 여백 밖 허용으로 한 줄이 되는 줄
    (전체 폭은 1 초과인데 앱이 줄바꿈 안 함으로 예측하는 줄), 없으면 1 초과 쪽 처음."""
    found = []
    for base in PUNCT_BASES:
        for n in range(1, len(base) + 1):
            text = base[:n].rstrip() + punct
            if BAND[0] <= raw_ratio(fit, text) <= BAND[1]:
                found.append(text)
    under = sorted((t for t in found if raw_ratio(fit, t) <= 1), key=lambda t: raw_ratio(fit, t))[-1:]
    hanging = sorted((t for t in found if raw_ratio(fit, t) > 1 and fit.ratio(t) <= 1), key=lambda t: raw_ratio(fit, t))[:1]
    over = sorted((t for t in found if fit.ratio(t) > 1), key=lambda t: raw_ratio(fit, t))[:1] if not hanging else []
    return under + hanging + over


def collect_cases(fit):
    cases, seen = [], set()
    for name, stream in STREAMS.items():
        for text in boundary_lines(fit, stream):
            if text not in seen:
                seen.add(text)
                cases.append((name, text))
    for p in PUNCTS:
        for text in punct_lines(fit, p):
            if text not in seen:
                seen.add(text)
                cases.append(('끝 문장부호 ' + p, text))
    return cases


def verdict(ratio):
    return '줄바꿈 예상' if ratio > 1 else '줄바꿈 안 함'


def add_label(slide, prs, text, color):
    box = slide.shapes.add_textbox(Emu(457200), prs.slide_height - Emu(640080), prs.slide_width - Emu(914400), Emu(457200))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(text.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(12)
        run.font.name = LABEL_FONT
        run.font.color.rgb = color


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_summary(prs, rows):
    """맨 앞 요약 슬라이드: 슬라이드 번호·예측 표 (가사 글상자가 없는 빈 레이아웃 위에 글상자로)"""
    slide = prs.slides.add_slide(prs.slides[0].slide_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    box = slide.shapes.add_textbox(Emu(228600), Emu(137160), prs.slide_width - Emu(457200), prs.slide_height - Emu(274320))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(rows):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(9)
        run.font.name = LABEL_FONT
    # 맨 앞으로 옮긴다
    id_list = prs.slides._sldIdLst
    id_list.insert(0, id_list[-1])
    return slide


def main():
    parser = argparse.ArgumentParser(description='Windows 실기 확인용 시험 PPT 만들기')
    parser.add_argument('-o', '--output', default='windows_check.pptx')
    args = parser.parse_args()

    app = QGuiApplication.instance() or QGuiApplication([])  # noqa: F841 (Qt 글꼴 API에 필요)
    font_dir = resource_path(os.path.join('assets', 'fonts'))
    for name in ('Pretendard-Regular.otf', 'Pretendard-SemiBold.otf', 'Pretendard-Bold.otf'):
        QFontDatabase.addApplicationFont(os.path.join(font_dir, name))

    template = resource_path('template.pptx')
    fit = LineFit(template)
    if not fit.box:
        sys.exit('템플릿에서 가사 글상자를 찾지 못했어요.')
    if not fit.box.exact:
        print('경고: 템플릿 글꼴(%s / %s)이 이 컴퓨터에 없어 예측이 어림값이에요.' % (fit.box.latin, fit.box.east_asian))

    cases = collect_cases(fit)
    multi = {
        3: ['첫째 줄이에요', '둘째 줄이에요', '셋째 줄이에요'],
        4: ['첫째 줄이에요', '둘째 줄이에요', '셋째 줄이에요', '넷째 줄이에요'],
        5: ['첫째 줄이에요', '둘째 줄이에요', '셋째 줄이에요', '넷째 줄이에요', '다섯째 줄이에요'],
    }
    chunks = [text for _, text in cases] + ['\n'.join(lines) for lines in multi.values()]

    builder = PPTBuilder(template)
    buffer = builder.generate_slides_to_memory(chunks, title='Windows 실기 확인')
    prs = Presentation(buffer)
    slides = list(prs.slides)

    summary = ['시험 PPT 요약 (앱 예측). 비율 = 글자 폭 / 글상자 쓸 수 있는 폭, 1 초과면 줄바꿈 예상',
               '템플릿 글상자: 폭 %.1fpt, 글꼴 %s / %s %.0fpt' % (
                   fit.box.available_pt, fit.box.latin, fit.box.east_asian, fit.box.size_pt),
               '(슬라이드 번호는 이 요약 슬라이드를 뺀 번호. 이 슬라이드가 1번이므로 실제로는 +1)']
    for i, (name, text) in enumerate(cases):
        ratio = fit.ratio(text)
        label = '[%d] %s  비율 %.4f  → %s' % (i + 1, name, ratio, verdict(ratio))
        if raw_ratio(fit, text) > 1 >= ratio:
            label += '  (전체 폭 %.4f, 줄 끝 문장부호는 여백 밖 허용으로 뺌)' % raw_ratio(fit, text)
        summary.append('%s   "%s"' % (label, text))
        add_label(slides[i], prs, label + '\n실제로 한 줄로 보이면 O, 두 줄이면 X로 기록', RGBColor(0xC0, 0x39, 0x2B) if ratio > 1 else RGBColor(0x1E, 0x7E, 0x34))
        set_notes(slides[i], '%s\n앱 예측 비율: %.4f → %s\n줄: %s' % (name, ratio, verdict(ratio), text))
    for k, lines in enumerate(multi):
        slide = slides[len(cases) + k]
        label = '[%d줄 슬라이드] 모든 줄이 보이는지, 첫 줄이 잘리지 않는지 확인' % lines
        add_label(slide, prs, label, RGBColor(0x33, 0x33, 0x33))
        set_notes(slide, label + '\n예측값 없음: 글상자 밖으로 넘치거나 첫 줄이 위로 잘리면 이상.')
        summary.append('[%d줄 슬라이드] 잘림 확인' % lines)

    add_summary(prs, summary)
    prs.save(args.output)
    print('%s 만듦 — 경계 줄 %d개 + 3·4·5줄 슬라이드 (요약 슬라이드 1장 포함 %d장)' % (
        args.output, len(cases), len(prs.slides)))
    for i, (name, text) in enumerate(cases):
        print('%3d  %-14s %.4f  %-8s %s' % (i + 1, name, fit.ratio(text), verdict(fit.ratio(text)), text))


if __name__ == '__main__':
    main()
